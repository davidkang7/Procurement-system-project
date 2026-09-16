# -*- coding: utf-8 -*-
"""
주문서 입력 소스 — 두 경로를 같은 PoOrder 로 변환한다.

  ① GAS 핸드오프(자동): PRC 최종승인 시 PO.gs 가 FINAL/{PO} 폴더에 쓴 po_manifest.json
  ② DB 백업(수동):     구매결재시스템_DB_YYYYMMDD.xlsx 의 '품의서목록' 탭 + 품의번호

② 는 결재 시스템을 거치지 않고 과거 건을 재생성하거나 핸드오프가 실패했을 때 쓰는 우회로다.
어느 쪽이든 **값을 생성·추론하지 않는다.** 없는 값은 비운다(문서 위조 방지).
"""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

from po_form import PoItem, PoOrder, split_unit_spec

SUPPORTED_SCHEMAS = ("po-manifest-v1", "po-manifest-v2")
DB_SHEET = "품의서목록"


def _to_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.match(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _to_num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    return float(s) if s not in ("", "-", ".") else 0.0


# ──────────────────────────────────────────────────────────────
# ① 매니페스트
# ──────────────────────────────────────────────────────────────
def from_manifest(m: dict) -> PoOrder:
    items = []
    currency = str(m.get("currency") or "KRW").upper()
    for it in (m.get("items") or []):
        unit = str(it.get("unit") or "").strip()
        spec = str(it.get("spec") or "").strip()
        if not unit:                       # v1 매니페스트: 둘째 필드가 unit/spec 혼용
            unit, spec = split_unit_spec(spec)
        items.append(PoItem(
            name=str(it.get("name") or "").strip(),
            qty=_to_num(it.get("qty")),
            price=_to_num(it.get("price")),
            unit=unit,
            spec=spec,
            currency=str(it.get("currency") or currency).upper(),
        ))
    if items:
        currency = items[0].currency

    return PoOrder(
        po_no=str(m.get("poNo") or "").strip(),
        vendor_name=str(m.get("vendorName") or "").strip(),
        items=items,
        issue_date=_to_date(m.get("issueDate")),
        delivery_date=_to_date(m.get("deliveryDate")),
        payment_terms_raw=str(m.get("paymentTerms") or "").strip(),
        destination=str(m.get("destinationOverride") or "").strip(),
        reference="",                       # Reference NO.(C11)는 관례상 공란
        total_amt=(_to_num(m.get("totalAmt")) if m.get("totalAmt") not in (None, "") else None),
        currency=currency,
        subject=str(m.get("subject") or "").strip(),
        prc_token=str(m.get("prcToken") or "").strip(),
    )


def load_manifest_file(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def check_schema(m: dict):
    schema = m.get("schemaVersion")
    if schema not in SUPPORTED_SCHEMAS:
        print(f"  · [경고] 매니페스트 스키마 불일치: {schema} (지원: {', '.join(SUPPORTED_SCHEMAS)})")


# ──────────────────────────────────────────────────────────────
# ② DB 백업 xlsx
# ──────────────────────────────────────────────────────────────
ITEM_PATTERNS = (
    # "1. 품명 | 규격/단위 | 30개 | 143,700 JPY"
    re.compile(r"^\d+\.\s*(?P<name>.+?)\s*\|\s*(?P<second>.*?)\s*\|\s*(?P<qty>[\d,]+)\s*개\s*\|"
               r"\s*(?P<price>[\d,]+(?:\.\d+)?)\s*(?P<cur>[A-Z]{3})$"),
    # "1. 품명 / 규격 / 30개 / 143,700 KRW"
    re.compile(r"^\d+\.\s*(?P<name>.+)\s*/\s*(?P<second>.+?)\s*/\s*(?P<qty>[\d,]+)\s*개\s*/"
               r"\s*(?P<price>[\d,]+(?:\.\d+)?)\s*(?P<cur>[A-Z]{3})$"),
    # 구 포맷: "1. 품명 규격 / 30개 / 143,700원"
    re.compile(r"^\d+\.\s*(?P<name>.+?)\s+(?P<second>.+?)\s*/\s*(?P<qty>[\d,]+)\s*개\s*/"
               r"\s*(?P<price>[\d,]+(?:\.\d+)?)\s*원$"),
)


def parse_items(text: str):
    """품목내역 텍스트 → [PoItem]. GAS parseItemsSummary 와 같은 규칙."""
    items = []
    for line in str(text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        for pat in ITEM_PATTERNS:
            m = pat.match(line)
            if not m:
                continue
            unit, spec = split_unit_spec(m.group("second"))
            cur = m.groupdict().get("cur") or "KRW"
            items.append(PoItem(
                name=m.group("name").strip(),
                qty=_to_num(m.group("qty")),
                price=_to_num(m.group("price")),
                unit=unit,
                spec=spec,
                currency=cur.upper(),
            ))
            break
        else:
            raise SystemExit(f"[오류] 품목내역을 해석하지 못했습니다: {line!r}\n"
                             "  DB 원문 형식을 확인하세요(품명 | 단위 | N개 | 단가 통화).")
    return items


def from_db(db_path: Path, po_no: str, require_approved: bool = True) -> PoOrder:
    import openpyxl

    wb = openpyxl.load_workbook(db_path, data_only=True, read_only=True)
    if DB_SHEET not in wb.sheetnames:
        raise SystemExit(f"[오류] '{DB_SHEET}' 시트가 없습니다: {db_path}")
    ws = wb[DB_SHEET]

    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    need = ["품의번호", "업체명", "품목내역", "합계금액", "DocType", "결재상태"]
    missing = [c for c in need if c not in idx]
    if missing:
        raise SystemExit(f"[오류] DB 컬럼 누락: {missing}")

    def val(row, name):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    target = None
    for row in rows:
        if str(val(row, "품의번호") or "").strip() == po_no and str(val(row, "DocType") or "") == "PRC":
            target = row                     # 같은 번호가 여러 행이면 마지막(최신) 행을 쓴다
    if target is None:
        raise SystemExit(f"[오류] DB에서 PRC 행을 찾지 못했습니다: {po_no}")

    status = str(val(target, "결재상태") or "")
    if require_approved and "최종승인(PRC)" not in status:
        raise SystemExit(f"[오류] 최종승인(PRC) 상태가 아닙니다: {po_no} / 현재 '{status}'\n"
                         "  결재 완료 전 주문서를 만들려면 --allow-unapproved 를 붙이세요.")

    items = parse_items(val(target, "품목내역"))
    currency = items[0].currency if items else "KRW"
    payment_raw = str(val(target, "구매조건") or val(target, "구매방법") or "").strip()

    return PoOrder(
        po_no=po_no,
        vendor_name=str(val(target, "업체명") or "").strip(),
        items=items,
        issue_date=_to_date(val(target, "발행일자")),
        delivery_date=_to_date(val(target, "납기일")),
        payment_terms_raw=payment_raw,
        total_amt=_to_num(val(target, "합계금액")),
        currency=currency,
        subject=str(val(target, "품의제목") or "").strip(),
        prc_token=str(val(target, "토큰") or "").strip(),
    )
